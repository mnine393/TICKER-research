"""Formatted Excel workbook export (openpyxl)."""
from __future__ import annotations

import io
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .checks import summary as check_summary

HEADER_FILL = PatternFill("solid", fgColor="1F3B5C")
TITLE_FONT = Font(bold=True, size=13, color="1F3B5C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
PASS_FILL = PatternFill("solid", fgColor="D9F2E6")
FAIL_FILL = PatternFill("solid", fgColor="F8D7DA")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT_WORDS = ("margin", "%", "growth", "rate", "upside", "share of", "yield", "vs. market", "cost of equity", "tax")
TOTAL_WORDS = ("total", "ebitda", "ebit", "net income", "fcfe", "ending", "equity value", "cash from")


def _is_pct_label(label: str) -> bool:
    l = str(label).lower()
    return any(w in l for w in PCT_WORDS) and "($" not in l and "($m" not in l


def _autowidth(ws, min_w=10, max_w=70):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=min_w)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, width + 2))


def _title(ws, row: int, text: str) -> int:
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
    return row + 1


def _write_table(ws, df: pd.DataFrame, row: int, title: str | None = None, index: bool = False,
                 pct_rows_by_label: bool = False, col_formats: dict | None = None, note: str | None = None) -> int:
    if title:
        row = _title(ws, row, title)
    if note:
        ws.cell(row=row, column=1, value=note).font = Font(italic=True, color="666666")
        row += 1
    data = df.reset_index() if index else df
    for j, c in enumerate(data.columns, start=1):
        cell = ws.cell(row=row, column=j, value=str(c))
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    row += 1
    col_formats = col_formats or {}
    for rec in data.itertuples(index=False):
        label = rec[0]
        is_pct_row = pct_rows_by_label and _is_pct_label(label)
        for j, v in enumerate(rec, start=1):
            if isinstance(v, (float, np.floating)) and np.isnan(v):
                v = None
            if isinstance(v, (np.integer, np.floating)):
                v = v.item()
            if isinstance(v, (list, dict, tuple)):
                v = str(v)
            cell = ws.cell(row=row, column=j, value=v)
            colname = data.columns[j - 1]
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                if colname in col_formats:
                    cell.number_format = col_formats[colname]
                elif is_pct_row:
                    cell.number_format = PCT
                else:
                    cell.number_format = NUM2 if abs(v) < 100 and v != int(v) else NUM
        if isinstance(label, str) and any(w in label.lower() for w in TOTAL_WORDS) and pct_rows_by_label:
            for j in range(1, len(data.columns) + 1):
                ws.cell(row=row, column=j).font = BOLD
                ws.cell(row=row, column=j).border = THIN
        row += 1
    return row + 1


def build_workbook(bundle, aset, scenario: str, evals: dict, extras: dict, source_table: pd.DataFrame,
                   hist_table: pd.DataFrame) -> bytes:
    """``extras`` holds sensitivity outputs: ke_g, tornado, one_var, two_var, bridge, wmbt, summary, multiples."""
    wb = Workbook()
    ev = evals[scenario]
    res, val = ev.model, ev.valuation

    # ---------------- Cover ----------------
    ws = wb.active
    ws.title = "Cover"
    r = _title(ws, 1, "Tesla, Inc. (TSLA) - Pro-forma FCFE valuation model (educational)")
    lines = [
        f"Generated: {datetime.now().isoformat(timespec='minutes')}",
        f"Scenario exported in detail: {config.SCENARIO_LABELS[scenario]}",
        f"Units: USD millions unless noted; per-share figures in USD; shares in millions.",
        f"Market price: ${bundle.market['price']['value']:,.2f} as of {bundle.market['price']['as_of']}",
        f"Implied value per share ({config.SCENARIO_LABELS[scenario]}): "
        + (f"${val.value_per_share:,.2f}" if val.status == "OK" else "NOT AVAILABLE - model checks failed"),
        "",
        "DISCLAIMER: " + config.DISCLAIMER,
        "",
        "Data status this session:",
    ] + [f"  [{s[1]}] {s[0]}: {s[2]}" for s in bundle.status]
    for ln in lines:
        ws.cell(row=r, column=1, value=ln).alignment = Alignment(wrap_text=True)
        r += 1
    ws.column_dimensions["A"].width = 140

    # ---------------- Sources ----------------
    ws = wb.create_sheet("Sources")
    _write_table(ws, source_table.drop(columns=["URL"]).assign(URL=source_table["URL"]), 1,
                 "Source table - every historical line item, obligation and market input")
    ws.freeze_panes = "C3"
    _autowidth(ws, max_w=60)

    # ---------------- Assumptions ----------------
    ws = wb.create_sheet("Assumptions")
    r = 1
    for s in config.SCENARIOS:
        t = aset.table(s)
        fmts = {c: "#,##0.000" for c in t.columns if c.startswith("FY")}
        r = _write_table(ws, t, r, f"Operating drivers - {config.SCENARIO_LABELS[s]} scenario", col_formats=fmts,
                         note="Types: history = held at/derived from reported data; guidance = company-stated; judgment = analyst estimate with rationale. Percent drivers are decimals (0.21 = 21%).")
    r = _write_table(ws, aset.params_table(), r, "Valuation and structural parameters")
    rev = aset.revisions_table()
    _write_table(ws, rev if not rev.empty else pd.DataFrame([{"timestamp": "", "item": "No edits this session", "scenario": "", "year": "", "old": "", "new": "", "note": ""}]),
                 r, "Revision history (edits made in the app this session)")
    ws.freeze_panes = "D3"
    _autowidth(ws, max_w=55)

    # ---------------- Historical Financials ----------------
    ws = wb.create_sheet("Historical Financials")
    _write_table(ws, hist_table, 1, "Reported historicals (USD m; KPIs in units/GWh) - FY2023A-FY2025A, H1 2026A, Q2 2026 balance sheet")
    ws.freeze_panes = "D3"
    _autowidth(ws, max_w=60)

    # ---------------- Forecast ----------------
    ws = wb.create_sheet("Forecast")
    r = _write_table(ws, res.income, 1, f"Income statement ({config.SCENARIO_LABELS[scenario]})", index=True, pct_rows_by_label=True)
    r = _write_table(ws, res.balance, r, "Balance sheet", index=True, pct_rows_by_label=True)
    r = _write_table(ws, res.cashflow, r, "Cash flow statement", index=True, pct_rows_by_label=True,
                     note="Historical CFO includes other non-cash items; historical CFI includes net purchases of short-term investments (forecast treats them as cash).")
    for name, df in res.schedules.items():
        r = _write_table(ws, df, r, name, index=True, pct_rows_by_label=True)
    _write_table(ws, res.fcfe, r, "FCFE build", index=True, pct_rows_by_label=True)
    ws.freeze_panes = "B2"
    _autowidth(ws, max_w=60)

    # ---------------- Valuation ----------------
    ws = wb.create_sheet("Valuation")
    r = 1
    if val.status != "OK":
        r = _title(ws, r, "VALUATION REFUSED - material model checks failed:")
        for f in val.failures:
            ws.cell(row=r, column=1, value=f).fill = FAIL_FILL
            r += 1
        r += 1
    coe = val.coe
    capm_df = pd.DataFrame([
        {"Item": "Risk-free rate", "Value": coe["rf"]}, {"Item": "Equity risk premium", "Value": coe["erp"]},
        {"Item": "Observed beta (regression)", "Value": coe["beta_observed"]},
        {"Item": "Bottom-up beta (relevered)", "Value": coe["beta_bottom_up"]},
        {"Item": f"Beta used ({coe['method']})", "Value": coe["beta"]},
        {"Item": "Cost of equity used (CAPM rf + beta x ERP)", "Value": val.ke},
        {"Item": "Cost of equity @ observed beta", "Value": coe["ke_observed"]},
        {"Item": "Cost of equity @ bottom-up beta", "Value": coe["ke_bottom_up"]},
        {"Item": "Terminal growth", "Value": val.g},
        {"Item": "Tesla forecast D/E used for relevering", "Value": coe["bottom_up"]["tesla_de"]},
    ])
    fm = {"Value": '0.000'}
    r = _write_table(ws, capm_df, r, "Cost of equity (CAPM)", col_formats=fm)
    for rr in range(r - len(capm_df) - 1, r - 1):
        label = ws.cell(row=rr, column=1).value or ""
        if any(w in label.lower() for w in ("rate", "premium", "cost of equity", "growth", "d/e")):
            ws.cell(row=rr, column=2).number_format = "0.00%"
    r = _write_table(ws, coe["bottom_up"]["peers"], r, "Bottom-up beta peers (Hamada unlevering)",
                     col_formats={"Levered beta": "0.00", "D/E": "0.00", "Tax rate": "0%", "Unlevered beta": "0.00"})
    bu = coe["bottom_up"]
    r = _write_table(ws, pd.DataFrame([{"Auto peers avg unlevered": bu["auto_unlevered"], "Energy peers avg unlevered": bu["energy_unlevered"],
                                        "Revenue-weighted unlevered": bu["weighted_unlevered"], "Relevered at Tesla D/E": bu["relevered"]}]),
                     r, "Bottom-up beta build", col_formats={k: "0.000" for k in ["Auto peers avg unlevered", "Energy peers avg unlevered", "Revenue-weighted unlevered", "Relevered at Tesla D/E"]})
    if val.status == "OK":
        r = _write_table(ws, val.pv_table, r, "Discounted operating FCFE", col_formats={"Discount factor": "0.0000", "Discount period (yrs)": "0.00"})
        r = _write_table(ws, val.terminal_bridge, r, "2031 terminal FCFE bridge")
        eb = val.equity_bridge.copy()
        r = _write_table(ws, eb, r, "Equity value bridge")
        for rr in range(r - len(eb) - 1, r - 1):
            label = (ws.cell(row=rr, column=1).value or "").lower()
            if "upside" in label or "share of" in label:
                ws.cell(row=rr, column=2).number_format = "0.0%"
            elif "per share" in label or "price" in label:
                ws.cell(row=rr, column=2).number_format = "$#,##0.00"
        if val.warnings:
            r = _title(ws, r, "Warnings")
            for w_ in val.warnings:
                ws.cell(row=r, column=1, value=w_)
                r += 1
            r += 1
    if "multiples" in extras:
        _write_table(ws, extras["multiples"], r, "SECONDARY cross-check: trading multiples (not the primary valuation)",
                     col_formats={"Multiple": "0.0x", "Implied value per share": "$#,##0.00"})
    _autowidth(ws, max_w=70)

    # ---------------- Scenarios ----------------
    ws = wb.create_sheet("Scenarios")
    r = _write_table(ws, extras["summary"], 1, "Scenario summary (each scenario = full model re-run)",
                     col_formats={"Upside / (downside)": PCT, "Cost of equity": PCT, "Terminal growth": PCT,
                                  "TV % of DCF value": PCT, "FY2030E EBITDA margin": PCT,
                                  "Implied value / share ($)": "$#,##0.00", "Market price ($)": "$#,##0.00"})
    if "bridge" in extras:
        r = _write_table(ws, extras["bridge"], r, "Price / implied-value bridge",
                         col_formats={"Value per share ($)": "$#,##0.00", "Cost of equity": PCT, "vs. market price": PCT})
    if "ke_g" not in extras:
        ws.cell(row=r, column=1, value="Sensitivities not computed: the valuation for this scenario was refused because a "
                                       "material model check failed (see the Checks sheet).").fill = FAIL_FILL
    else:
        r = _write_table(ws, extras["ke_g"], r, "Cost of equity x terminal growth ($/share; blank = g >= ke, valuation refused)",
                         index=True, col_formats={c: "$#,##0.00" for c in extras["ke_g"].columns})
        tor = extras["tornado"].drop(columns=["ID"])
        r = _write_table(ws, tor, r, "Tornado (ranked by range; full model re-run for each shock)",
                         col_formats={c: "$#,##0.00" for c in tor.columns if "value" in c.lower() or "delta" in c.lower() or c == "Range"})
        ov = extras["one_var"]
        r = _write_table(ws, ov, r, "One-variable sensitivities ($/share at each shock step)",
                         col_formats={c: "$#,##0.00" for c in ov.columns if c.startswith("step")})
        r = _write_table(ws, extras["two_var"], r, "Two-variable matrix - two highest-impact operating drivers ($/share)", index=True,
                         col_formats={c: "$#,##0.00" for c in extras["two_var"].columns})
        _write_table(ws, extras["wmbt"], r, "What would have to be true for today's market price to be correct?")
    _autowidth(ws, max_w=60)

    # ---------------- Checks ----------------
    ws = wb.create_sheet("Checks")
    chk = val.checks.copy()
    r = _write_table(ws, check_summary(chk), 1, "Model integrity checks - summary")
    start = r + 1
    _write_table(ws, chk, r, "Model integrity checks - every forecast year", col_formats={"Difference": "#,##0.000"})
    for i, passed in enumerate(chk["Pass"]):
        fill = PASS_FILL if passed else FAIL_FILL
        for j in range(1, len(chk.columns) + 1):
            ws.cell(row=start + 1 + i, column=j).fill = fill
    _autowidth(ws, max_w=80)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
