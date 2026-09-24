"""Scenario/sensitivity engine (full re-runs, no additivity) and the Excel export."""
import io

import numpy as np
import openpyxl
import pytest

from tsla_model import config
from tsla_model import scenarios as S
from tsla_model.data_ingestion import build_source_table, historical_table
from tsla_model.excel_export import build_workbook


def test_scenarios_are_ordered_and_valid(bundle, aset, val_date):
    ev = S.run_all_scenarios(bundle, aset, val_date)
    vals = {s: e.value for s, e in ev.items()}
    assert all(e.valuation.status == "OK" for e in ev.values())
    assert vals["bull"] > vals["base"] > vals["bear"]
    assert vals["base"] > vals["recession"]


def test_two_variable_matrix_is_full_rerun_not_additive(bundle, aset, val_date):
    tor = S.tornado(bundle, aset, "base", val_date)
    x, y = tor["ID"][0], tor["ID"][1]
    m = S.two_variable(bundle, aset, x, y, "base", val_date)
    base = S.evaluate(bundle, aset, "base", val_date).value
    assert m.iloc[2, 2] == pytest.approx(base)
    corner = m.iloc[4, 4]
    direct = S.evaluate(bundle, aset.shocked(y, "base", 2).shocked(x, "base", 2), "base", val_date).value
    assert corner == pytest.approx(direct)
    dx = m.iloc[2, 4] - base
    dy = m.iloc[4, 2] - base
    assert corner != pytest.approx(base + dx + dy, abs=1e-6)   # interaction captured => not simply added


def test_tornado_sorted_and_covers_core_drivers(bundle, aset, val_date):
    tor = S.tornado(bundle, aset, "base", val_date)
    assert list(tor["Range"]) == sorted(tor["Range"], reverse=True)
    assert set(tor["ID"]) == set(aset.core_drivers())
    ov = S.one_variable(bundle, aset, "base", val_date)
    assert set(ov["ID"]) == set(aset.core_drivers())
    base = S.evaluate(bundle, aset, "base", val_date).value
    assert np.allclose(ov["step +0"], base)


def test_ke_g_grid_blanks_invalid_cells(bundle, aset, val_date):
    grid = S.ke_g_grid(bundle, aset, "base", val_date, ke_values=[0.05, 0.10], g_values=[0.03, 0.06])
    assert np.isnan(grid.loc["5.00%", "6.00%"])
    assert not np.isnan(grid.loc["10.00%", "3.00%"])
    assert grid.loc["10.00%", "6.00%"] > grid.loc["10.00%", "3.00%"]


def test_what_must_be_true_solutions_reprice_to_market(bundle, aset, val_date):
    w = S.what_must_be_true(bundle, aset, "base", val_date)
    col = [c for c in w.columns if c.startswith("Required")][0]
    ke_req = float(w.loc[w.iloc[:, 0] == "Cost of equity", col].iloc[0].rstrip("%")) / 100
    v = S.evaluate(bundle, aset, "base", val_date, ke_override=ke_req).value
    assert v == pytest.approx(bundle.market["price"]["value"], rel=0.01)


def test_value_bridge_contains_required_cases(bundle, aset, val_date):
    ev = S.run_all_scenarios(bundle, aset, val_date)
    br = S.value_bridge(bundle, aset, ev, val_date)
    cases = " ".join(br["Case"])
    for k in ["Market price", "Bear", "Base case", "Bull", "observed beta", "cost of equity"]:
        assert k in cases


def test_excel_export_has_required_sheets(bundle, aset, val_date):
    x = S.full_analysis(bundle, aset, "base", val_date)
    data = build_workbook(bundle, aset, "base", x["evals"], x, build_source_table(bundle), historical_table(bundle))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for name in ["Sources", "Assumptions", "Historical Financials", "Forecast", "Valuation", "Scenarios", "Checks"]:
        assert name in wb.sheetnames
    assert "not investment advice" in " ".join(str(c.value) for c in wb["Cover"]["A"] if c.value).lower()
    assert wb["Sources"].max_row > 300
